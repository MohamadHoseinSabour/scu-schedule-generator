import openpyxl
from pathlib import Path

class XlsxSheetWrapper:
    def __init__(self, ws):
        self.ws = ws

    def cell_value(self, rowx: int, colx: int):
        # xlrd is 0-indexed, openpyxl is 1-indexed
        val = self.ws.cell(row=rowx + 1, column=colx + 1).value
        if val is None:
            return ""
        return val

class XlsxBookWrapper:
    def __init__(self, wb):
        self.wb = wb

    @property
    def nsheets(self) -> int:
        return len(self.wb.sheetnames)

    def sheet_by_index(self, index: int) -> XlsxSheetWrapper:
        ws = self.wb.worksheets[index]
        return XlsxSheetWrapper(ws)

def read_xlsx(file_path: Path) -> XlsxBookWrapper:
    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    return XlsxBookWrapper(wb)
